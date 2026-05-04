from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/home/xxs/myboke")
TARGET_DIR = ROOT / "static" / "tools" / "yysls-graduation"
SOURCES = [
    {
        "name": "破竹鸢",
        "source": ROOT / "yysls" / "破竹鸢100级竞速轴属性毕业率计算器4.3.xlsx",
        "target": TARGET_DIR / "workbook-yuan.json",
    },
    {
        "name": "破竹尘",
        "source": ROOT / "yysls" / "破竹尘100级竞速轴属性毕业率计算器4.5.xlsx",
        "target": TARGET_DIR / "workbook-chen.json",
    },
]


def cell_payload(value):
    if value is None:
        return None
    if isinstance(value, str) and value.startswith("="):
        return {"f": value}
    return {"v": value}


def extract_workbook(source: Path, name: str) -> dict:
    wb = load_workbook(source, data_only=False, read_only=False)
    payload = {
        "meta": {
            "name": name,
            "filename": source.name,
        },
        "sheetOrder": wb.sheetnames,
        "sheets": {},
    }

    for ws in wb.worksheets:
        sheet_cells = {}
        min_row = None
        max_row = 0
        min_col = None
        max_col = 0

        for row in ws.iter_rows():
            for cell in row:
                data = cell_payload(cell.value)
                if data is None:
                    continue
                sheet_cells[cell.coordinate] = data
                min_row = cell.row if min_row is None else min(min_row, cell.row)
                min_col = cell.column if min_col is None else min(min_col, cell.column)
                max_row = max(max_row, cell.row)
                max_col = max(max_col, cell.column)

        payload["sheets"][ws.title] = {
            "cells": sheet_cells,
            "bounds": {
                "minRow": min_row or 1,
                "maxRow": max_row,
                "minCol": min_col or 1,
                "maxCol": max_col,
            },
        }

    return payload


def main() -> None:
    TARGET_DIR.mkdir(parents=True, exist_ok=True)

    for item in SOURCES:
        payload = extract_workbook(item["source"], item["name"])
        item["target"].write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))
        print(f"Wrote {item['target']}")


if __name__ == "__main__":
    main()
