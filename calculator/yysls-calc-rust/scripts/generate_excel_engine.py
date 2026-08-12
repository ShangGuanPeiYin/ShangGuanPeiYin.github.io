#!/usr/bin/env python3
from __future__ import annotations

import json
import copy
import gc
import math
import os
import re
import struct
from dataclasses import dataclass
from datetime import timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string


ROOT = Path(__file__).resolve().parents[3]
EXCEL_DIR = ROOT / "static/tools/yysls-tiaolv/excels"
METADATA_PATH = ROOT / "static/tools/yysls-tiaolv/assets/js/generated-calc-metadata.js"
STRINGS_PATH = ROOT / "static/tools/yysls-tiaolv/assets/js/generated-calc-strings.js"
DIRECT_SOURCE_DIR = ROOT / "calculator/yysls-calc-rust/target/generated-excel-direct"
MANIFEST_PATH = ROOT / "calculator/yysls-calc-rust/excel-sources.json"

SUPPORTED_FUNCTIONS = {"IF", "IFERROR", "VLOOKUP", "XLOOKUP", "OR", "MIN", "MAX", "SUM"}
FLOW_ORDER = ["牵丝玉", "牵丝翊", "破竹尘", "破竹风", "破竹鸢", "裂石威", "裂石钧", "鸣金虹", "鸣金影", "牵丝霖"]
FLOW_SLUGS = {"牵丝玉": "qsyu", "牵丝翊": "qsyi", "破竹尘": "pzchen", "破竹风": "pzfeng", "破竹鸢": "pzyuan", "裂石威": "lswei", "裂石钧": "lsjun", "鸣金虹": "mjhong", "鸣金影": "mjying", "牵丝霖": "qslin"}
INPUT_LEN = 40


def load_js_assignment(path: Path, prefix: str):
    text = path.read_text(encoding="utf-8")
    match = re.search(re.escape(prefix) + r"(.*?);\s*(?:\n|$)", text, re.S)
    if not match:
        raise ValueError(f"无法解析 {path}")
    return json.loads(match.group(1))


metadata = load_js_assignment(METADATA_PATH, "window.YYSLS_CALC_METADATA=")
source_metadata = copy.deepcopy(metadata)
strings = load_js_assignment(STRINGS_PATH, "window.YYSLS_CALC_STRINGS=")
string_ids = {value: index for index, value in enumerate(strings)}


def string_id(value: str) -> int:
    if value not in string_ids:
        string_ids[value] = len(strings)
        strings.append(value)
    return string_ids[value]


def f64_literal(value: float) -> str:
    if math.isnan(value):
        return "f64::NAN"
    if math.isinf(value):
        return "f64::INFINITY" if value > 0 else "f64::NEG_INFINITY"
    bits = struct.unpack("<Q", struct.pack("<d", float(value)))[0]
    return f"f64::from_bits(0x{bits:016x})"


@dataclass
class Node:
    kind: str
    value: object = None
    args: list["Node"] | None = None


class FormulaParser:
    def __init__(self, formula: str):
        self.tokens = Tokenizer(formula).items
        self.pos = 0

    def peek(self):
        return self.tokens[self.pos] if self.pos < len(self.tokens) else None

    def take(self):
        token = self.peek()
        if token is None:
            raise ValueError("公式意外结束")
        self.pos += 1
        return token

    def parse(self) -> Node:
        result = self.parse_comparison()
        if self.pos != len(self.tokens):
            raise ValueError(f"公式尾部存在未解析内容：{self.tokens[self.pos:]}")
        return result

    def parse_comparison(self):
        node = self.parse_additive()
        while self.peek() and self.peek().type == "OPERATOR-INFIX" and self.peek().value in ("=", "<", ">=", ">", "<=", "<>"):
            op = self.take().value
            node = Node("binary", op, [node, self.parse_additive()])
        return node

    def parse_additive(self):
        node = self.parse_multiplicative()
        while self.peek() and self.peek().type == "OPERATOR-INFIX" and self.peek().value in ("+", "-"):
            op = self.take().value
            node = Node("binary", op, [node, self.parse_multiplicative()])
        return node

    def parse_multiplicative(self):
        node = self.parse_atom()
        while self.peek() and self.peek().type == "OPERATOR-INFIX" and self.peek().value in ("*", "/"):
            op = self.take().value
            node = Node("binary", op, [node, self.parse_atom()])
        return node

    def parse_atom(self):
        token = self.take()
        if token.type == "PAREN" and token.subtype == "OPEN":
            node = self.parse_comparison()
            close = self.take()
            if close.type != "PAREN" or close.subtype != "CLOSE":
                raise ValueError("括号不匹配")
        elif token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value[:-1].upper().removeprefix("_XLFN.")
            if name not in SUPPORTED_FUNCTIONS:
                raise ValueError(f"不支持的函数：{name}")
            args = []
            if not (self.peek() and self.peek().type == "FUNC" and self.peek().subtype == "CLOSE"):
                while True:
                    args.append(self.parse_comparison())
                    if self.peek() and self.peek().type == "SEP" and self.peek().subtype == "ARG":
                        self.take()
                        continue
                    break
            close = self.take()
            if close.type != "FUNC" or close.subtype != "CLOSE":
                raise ValueError(f"函数 {name} 缺少右括号")
            node = Node("func", name, args)
        elif token.type == "OPERAND":
            if token.subtype == "NUMBER":
                node = Node("number", float(token.value))
            elif token.subtype == "TEXT":
                node = Node("text", token.value[1:-1].replace('""', '"'))
            elif token.subtype == "LOGICAL":
                node = Node("number", 1.0 if token.value.upper() == "TRUE" else 0.0)
            elif token.subtype == "RANGE":
                node = Node("range", token.value)
            else:
                raise ValueError(f"不支持的操作数：{token.value}/{token.subtype}")
        else:
            raise ValueError(f"不支持的公式标记：{token.value}/{token.type}/{token.subtype}")
        while self.peek() and self.peek().type == "OPERATOR-POSTFIX" and self.peek().value == "%":
            self.take()
            node = Node("binary", "/", [node, Node("number", 100.0)])
        return node


class WorkbookCompiler:
    def __init__(self, path: Path, flow_name: str, class_name: str, input_cells: list[str], input_kinds: list[str]):
        self.path = path
        self.direct_mode = False
        self.flow_name = flow_name
        self.class_name = class_name
        self.workbook = load_workbook(path, data_only=False, read_only=False)
        self.values = load_workbook(path, data_only=True, read_only=True)
        self.sheets = {ws.title: ws for ws in self.workbook.worksheets}
        self.sheet_values = {ws.title: ws for ws in self.values.worksheets}
        self.cells: dict[tuple[str, str], int] = {}
        self.cell_names: list[str] = []
        self.cell_values: list[object] = []
        for ws in self.workbook.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        self.cells[(ws.title, cell.coordinate)] = len(self.cell_values)
                        self.cell_names.append(f"{ws.title}!{cell.coordinate}")
                        self.cell_values.append(cell.value.text if isinstance(cell.value, ArrayFormula) else cell.value)
        self.input_ids = [self.resolve_single(cell, "期望") for cell in input_cells]
        self.input_kinds = input_kinds
        wanted = set(input_cells) | {"期望!I8", "期望!I10", "期望!I12", "期望!I14", "RD!I14"}
        self.cached_values = {}
        for reference in wanted:
            sheet, address = self.split_ref(reference, "期望")
            self.cached_values[f"{sheet}!{address}"] = self.sheet_values[sheet][address].value
        self.values.close()
        self.sheet_values = {}
        self.lookup_rows = []
        lookup = self.sheets["武学奇术"]
        for row in range(1, lookup.max_row + 1):
            key = lookup.cell(row, 1).value
            if isinstance(key, str) and key:
                self.lookup_rows.append((string_id(key), row))
        self.gain_rows = []
        gain = self.sheets["增益"]
        for row in range(1, gain.max_row + 1):
            key = gain.cell(row, 1).value
            if isinstance(key, str) and key:
                self.gain_rows.append((string_id(key), row))
        self._const_ids = set()
        self._const_cache = {}
        self._pairs_key_cells = set()

    def split_ref(self, reference: str, current_sheet: str):
        cleaned = reference.replace("$", "")
        if "!" in cleaned:
            sheet, address = cleaned.rsplit("!", 1)
            sheet = sheet.strip("'").replace("''", "'")
        else:
            sheet, address = current_sheet, cleaned
        return sheet, address

    def resolve_single(self, reference: str, current_sheet: str) -> int | None:
        sheet, address = self.split_ref(reference, current_sheet)
        return self.cells.get((sheet, address))

    def range_ids(self, reference: str, current_sheet: str):
        sheet, address = self.split_ref(reference, current_sheet)
        ws = self.sheets[sheet]
        left, right = address.split(":", 1)
        if re.fullmatch(r"[A-Z]+", left) and re.fullmatch(r"[A-Z]+", right):
            min_col, max_col = column_index_from_string(left), column_index_from_string(right)
            min_row, max_row = 1, ws.max_row
        else:
            lcol, lrow = coordinate_from_string(left)
            rcol, rrow = coordinate_from_string(right)
            min_col, max_col = column_index_from_string(lcol), column_index_from_string(rcol)
            min_row, max_row = lrow, rrow
        result = []
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                cell_id = self.cells.get((sheet, ws.cell(row, col).coordinate))
                if cell_id is not None:
                    result.append(cell_id)
        return result

    def expr(self, node: Node, current_sheet: str) -> str:
        if node.kind == "number":
            return f64_literal(node.value)
        if node.kind == "text":
            return f64_literal(float(string_id(node.value)))
        if node.kind == "range":
            if ":" in str(node.value):
                raise ValueError(f"范围只能作为函数参数：{node.value}")
            cell_id = self.resolve_single(str(node.value), current_sheet)
            return f64_literal(0.0) if cell_id is None else (f"v[{cell_id}]" if self.direct_mode else f"self.cell({cell_id})")
        if node.kind == "binary":
            left = self.expr(node.args[0], current_sheet)
            right = self.expr(node.args[1], current_sheet)
            if node.value in ("+", "-", "*", "/"):
                return f"({left} {node.value} {right})"
            cmp = {"=": "==", "<>": "!=", "<": "<", ">": ">", "<=": "<=", ">=": ">="}[node.value]
            return f"if {left} {cmp} {right} {{ 1.0 }} else {{ 0.0 }}"
        if node.kind == "func":
            name = node.value
            if name == "IF":
                if len(node.args) != 3:
                    raise ValueError("IF参数数量必须为3")
                return f"if {self.expr(node.args[0], current_sheet)} != 0.0 {{ {self.expr(node.args[1], current_sheet)} }} else {{ {self.expr(node.args[2], current_sheet)} }}"
            if name == "OR":
                return "if " + " || ".join(f"{self.expr(arg, current_sheet)} != 0.0" for arg in node.args) + " { 1.0 } else { 0.0 }"
            if name == "IFERROR":
                if len(node.args) != 2:
                    raise ValueError("IFERROR参数数量必须为2")
                return self.expr(node.args[0], current_sheet)
            if name in ("MIN", "MAX"):
                method = "min" if name == "MIN" else "max"
                parts = [self.expr(arg, current_sheet) for arg in node.args]
                result = parts[0]
                for part in parts[1:]:
                    result = f"({result}).{method}({part})"
                return result
            if name == "SUM":
                parts = []
                for arg in node.args:
                    if arg.kind == "func" and arg.value == "XLOOKUP":
                        lookup_value, lookup_array, return_array, default = arg.args
                        lookup_sheet, lookup_address = self.split_ref(str(lookup_array.value), current_sheet)
                        return_sheet, return_address = self.split_ref(str(return_array.value), current_sheet)
                        if lookup_value.kind != "range" or lookup_sheet != "增益" or return_sheet != "增益" or lookup_address.replace("$", "") != "A:A":
                            raise ValueError("仅支持增益表范围XLOOKUP")
                        return_col = column_index_from_string(return_address.replace("$", "").split(":", 1)[0])
                        fallback = self.expr(default, current_sheet)
                        for cell_id in self.range_ids(str(lookup_value.value), current_sheet):
                            folded = self._try_fold_gain(cell_id, return_col, fallback)
                            parts.append(folded if folded is not None else (f"xlookup_gain(v[{cell_id}], {return_col}, {fallback}, v)" if self.direct_mode else f"self.xlookup_gain_cell({cell_id}, {return_col}, {fallback})"))
                        continue
                    if arg.kind == "range" and ":" in str(arg.value):
                        for cell_id in self.range_ids(str(arg.value), current_sheet):
                            value = self.cell_values[cell_id]
                            if isinstance(value, (int, float)) and not isinstance(value, bool) or isinstance(value, str) and value.startswith("="):
                                parts.append(f"v[{cell_id}]" if self.direct_mode else f"self.cell({cell_id})")
                    else:
                        parts.append(self.expr(arg, current_sheet))
                return "(" + " + ".join(parts or ["0.0"]) + ")"
            if name == "VLOOKUP":
                if len(node.args) != 4 or node.args[1].kind != "range" or node.args[2].kind != "number" or node.args[3].kind != "number" or node.args[3].value != 0:
                    raise ValueError("仅支持精确匹配VLOOKUP")
                key = self.expr(node.args[0], current_sheet)
                col = int(node.args[2].value)
                table_sheet, table_address = self.split_ref(str(node.args[1].value), current_sheet)
                if table_sheet == "武学奇术" and table_address.replace("$", "") == "A:ZZ":
                    if node.args[0].kind == "range" and ":" not in str(node.args[0].value):
                        key_id = self.resolve_single(str(node.args[0].value), current_sheet)
                        if key_id is None:
                            return f64_literal(0.0)
                        folded = self._try_fold_vlookup_global(key_id, col)
                        return folded if folded is not None else (f"vlookup(v[{key_id}], {col}, v)" if self.direct_mode else f"self.vlookup_cell({key_id}, {col})")
                    return f"vlookup({key}, {col}, v)" if self.direct_mode else f"self.vlookup({key}, {col})"
                left, right = table_address.replace("$", "").split(":", 1)
                lcol, lrow = coordinate_from_string(left)
                rcol, rrow = coordinate_from_string(right)
                first_col = column_index_from_string(lcol)
                pairs = []
                for row in range(lrow, rrow + 1):
                    key_id = self.cells.get((table_sheet, self.sheets[table_sheet].cell(row, first_col).coordinate))
                    value_id = self.cells.get((table_sheet, self.sheets[table_sheet].cell(row, first_col + col - 1).coordinate))
                    pairs.append(f"({key_id if key_id is not None else 'usize::MAX'}, {value_id if value_id is not None else 'usize::MAX'})")
                if node.args[0].kind == "range" and ":" not in str(node.args[0].value):
                    key_id = self.resolve_single(str(node.args[0].value), current_sheet)
                    if key_id is None:
                        return f64_literal(0.0)
                    folded = self._try_fold_local_vlookup(key_id, col, table_sheet, table_address)
                    return folded if folded is not None else (f"local_vlookup(v[{key_id}], &[{', '.join(pairs)}], v)" if self.direct_mode else f"self.local_vlookup_cell({key_id}, &[{', '.join(pairs)}])")
                return f"local_vlookup({key}, &[{', '.join(pairs)}], v)" if self.direct_mode else f"self.local_vlookup({key}, &[{', '.join(pairs)}])"
        raise ValueError(f"无法生成表达式：{node}")

    def cell_arm(self, cell_id: int):
        value = self.cell_values[cell_id]
        if isinstance(value, str) and value.startswith("="):
            parser = FormulaParser(value)
            node = parser.parse()
            sheet = next(sheet for (sheet, address), found in self.cells.items() if found == cell_id)
            return self.expr(node, sheet)
        if isinstance(value, bool):
            return f64_literal(1.0 if value else 0.0)
        if isinstance(value, (int, float)):
            return f64_literal(float(value))
        if isinstance(value, str):
            return f64_literal(float(string_id(value)))
        return f64_literal(0.0)

    def direct_dependencies(self):
        output_ids = [self.resolve_single(ref, "期望") for ref in ("I10", "I12", "I16", "I14")]
        reachable = set()
        global_lookup_cols = set()
        gain_lookup_cols = set()

        def add_ref(reference, current_sheet):
            if ":" in str(reference):
                for found in self.range_ids(str(reference), current_sheet):
                    visit(found)
            else:
                found = self.resolve_single(str(reference), current_sheet)
                if found is not None:
                    visit(found)

        def walk(node, current_sheet):
            if node.kind == "range":
                add_ref(node.value, current_sheet)
                return
            if node.kind != "func":
                for arg in node.args or []:
                    walk(arg, current_sheet)
                return
            if node.value == "VLOOKUP":
                walk(node.args[0], current_sheet)
                col = int(node.args[2].value)
                table_sheet, table_address = self.split_ref(str(node.args[1].value), current_sheet)
                if table_sheet == "武学奇术" and table_address.replace("$", "") == "A:ZZ":
                    global_lookup_cols.add(col)
                    for _, row in self.lookup_rows:
                        found = self.cells.get(("武学奇术", self.sheets["武学奇术"].cell(row, col).coordinate))
                        if found is not None:
                            visit(found)
                else:
                    left, right = table_address.replace("$", "").split(":", 1)
                    lcol, lrow = coordinate_from_string(left)
                    _, rrow = coordinate_from_string(right)
                    first_col = column_index_from_string(lcol)
                    for row in range(lrow, rrow + 1):
                        for offset in (0, col - 1):
                            found = self.cells.get((table_sheet, self.sheets[table_sheet].cell(row, first_col + offset).coordinate))
                            if found is not None:
                                visit(found)
                return
            if node.value == "XLOOKUP":
                lookup_value, _, return_array, default = node.args
                walk(lookup_value, current_sheet)
                walk(default, current_sheet)
                return_sheet, return_address = self.split_ref(str(return_array.value), current_sheet)
                col = column_index_from_string(return_address.replace("$", "").split(":", 1)[0])
                if return_sheet == "增益":
                    gain_lookup_cols.add(col)
                    for _, row in self.gain_rows:
                        found = self.cells.get(("增益", self.sheets["增益"].cell(row, col).coordinate))
                        if found is not None:
                            visit(found)
                return
            for arg in node.args or []:
                walk(arg, current_sheet)

        def visit(cell_id):
            if cell_id is None or cell_id in reachable:
                return
            reachable.add(cell_id)
            value = self.cell_values[cell_id]
            if isinstance(value, str) and value.startswith("="):
                sheet = self.cell_names[cell_id].split("!", 1)[0]
                walk(FormulaParser(value).parse(), sheet)

        for output_id in output_ids:
            visit(output_id)
        return reachable, global_lookup_cols, gain_lookup_cols

    def direct_order(self, reachable):
        """Return a stable dependency-first order for direct, single-evaluation code."""
        dependencies = {cell_id: set() for cell_id in reachable}

        def add_reference(target, reference, current_sheet):
            ids = self.range_ids(str(reference), current_sheet) if ":" in str(reference) else [self.resolve_single(str(reference), current_sheet)]
            dependencies[target].update(found for found in ids if found in reachable)

        def walk(target, node, current_sheet):
            if node.kind == "range":
                add_reference(target, node.value, current_sheet)
                return
            if node.kind != "func":
                for arg in node.args or []:
                    walk(target, arg, current_sheet)
                return
            if node.value == "VLOOKUP":
                walk(target, node.args[0], current_sheet)
                col = int(node.args[2].value)
                table_sheet, table_address = self.split_ref(str(node.args[1].value), current_sheet)
                if table_sheet == "武学奇术" and table_address.replace("$", "") == "A:ZZ":
                    for _, row in self.lookup_rows:
                        found = self.cells.get(("武学奇术", self.sheets["武学奇术"].cell(row, col).coordinate))
                        if found in reachable:
                            dependencies[target].add(found)
                else:
                    left, right = table_address.replace("$", "").split(":", 1)
                    lcol, lrow = coordinate_from_string(left)
                    _, rrow = coordinate_from_string(right)
                    first_col = column_index_from_string(lcol)
                    for row in range(lrow, rrow + 1):
                        for offset in (0, col - 1):
                            found = self.cells.get((table_sheet, self.sheets[table_sheet].cell(row, first_col + offset).coordinate))
                            if found in reachable:
                                dependencies[target].add(found)
                return
            if node.value == "XLOOKUP":
                walk(target, node.args[0], current_sheet)
                walk(target, node.args[3], current_sheet)
                return_sheet, return_address = self.split_ref(str(node.args[2].value), current_sheet)
                col = column_index_from_string(return_address.replace("$", "").split(":", 1)[0])
                if return_sheet == "增益":
                    for _, row in self.gain_rows:
                        found = self.cells.get(("增益", self.sheets["增益"].cell(row, col).coordinate))
                        if found in reachable:
                            dependencies[target].add(found)
                return
            for arg in node.args or []:
                walk(target, arg, current_sheet)

        for cell_id in reachable:
            value = self.cell_values[cell_id]
            if isinstance(value, str) and value.startswith("="):
                walk(cell_id, FormulaParser(value).parse(), self.cell_names[cell_id].split("!", 1)[0])

        order, state = [], {}
        def visit(cell_id):
            if state.get(cell_id) == 2:
                return
            if state.get(cell_id) == 1:
                raise ValueError(f"Excel公式循环引用：{self.cell_names[cell_id]}")
            state[cell_id] = 1
            for dependency in sorted(dependencies[cell_id]):
                visit(dependency)
            state[cell_id] = 2
            order.append(cell_id)
        for cell_id in sorted(reachable):
            visit(cell_id)
        return order, dependencies

    def direct_value_expr(self, cell_id):
        self.direct_mode = True
        try:
            return self.cell_arm(cell_id)
        finally:
            self.direct_mode = False

    def _const_value_cell(self, cell_id):
        if cell_id not in self._const_cache:
            raise ValueError(f"非常量单元格无法求值：{self.cell_names[cell_id]}")
        return self._const_cache[cell_id]

    def _const_value_of(self, cell_id):
        value = self.cell_values[cell_id]
        if isinstance(value, str) and value.startswith("="):
            parser = FormulaParser(value)
            node = parser.parse()
            sheet = next(sheet for (sheet, address), found in self.cells.items() if found == cell_id)
            return self._const_expr(node, sheet)
        return self._const_literal(value)

    def _cell_ref(self, cell_id):
        return f"v[{cell_id}]" if self.direct_mode else f"self.cell({cell_id})"

    def _try_fold_gain(self, cell_id: int, return_col: int, fallback: str):
        if cell_id not in self._const_ids:
            return None
        key = self._const_cache[cell_id]
        for key_id, row in self.gain_rows:
            if float(key_id) == key:
                value_id = self.cells.get(("增益", self.sheets["增益"].cell(row, return_col).coordinate))
                return fallback if value_id is None else self._cell_ref(value_id)
        return fallback

    def _try_fold_vlookup_global(self, key_id: int, col: int):
        if key_id not in self._const_ids:
            return None
        key = self._const_cache[key_id]
        for match_id, row in self.lookup_rows:
            if float(match_id) == key:
                value_id = self.cells.get(("武学奇术", self.sheets["武学奇术"].cell(row, col).coordinate))
                return f64_literal(0.0) if value_id is None else self._cell_ref(value_id)
        return f64_literal(0.0)

    def _try_fold_local_vlookup(self, key_id: int, col: int, table_sheet: str, table_address: str):
        if key_id not in self._const_ids:
            return None
        key = self._const_cache[key_id]
        left, right = table_address.replace("$", "").split(":", 1)
        lcol, lrow = coordinate_from_string(left)
        rcol, rrow = coordinate_from_string(right)
        first_col = column_index_from_string(lcol)
        for row in range(lrow, rrow + 1):
            kid = self.cells.get((table_sheet, self.sheets[table_sheet].cell(row, first_col).coordinate))
            vid = self.cells.get((table_sheet, self.sheets[table_sheet].cell(row, first_col + col - 1).coordinate))
            if kid is None:
                continue
            if kid not in self._const_ids:
                return None
            if self._const_cache[kid] == key:
                return f64_literal(0.0) if vid is None else self._cell_ref(vid)
        return f64_literal(0.0)

    def _const_literal(self, value) -> float:
        if isinstance(value, bool):
            return 1.0 if value else 0.0
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            return float(string_id(value))
        return 0.0

    def _const_expr(self, node: Node, current_sheet: str) -> float:
        if node.kind == "number":
            return float(node.value)
        if node.kind == "text":
            return float(string_id(node.value))
        if node.kind == "range":
            cell_id = self.resolve_single(str(node.value), current_sheet)
            return 0.0 if cell_id is None else self._const_value_cell(cell_id)
        if node.kind == "binary":
            left = self._const_expr(node.args[0], current_sheet)
            right = self._const_expr(node.args[1], current_sheet)
            if node.value == "+":
                return left + right
            if node.value == "-":
                return left - right
            if node.value == "*":
                return left * right
            if node.value == "/":
                return left / right
            cmp = {"=": left == right, "<>": left != right, "<": left < right, ">": left > right, "<=": left <= right, ">=": left >= right}[node.value]
            return 1.0 if cmp else 0.0
        if node.kind == "func":
            name = node.value
            if name == "IF":
                return self._const_expr(node.args[1], current_sheet) if self._const_expr(node.args[0], current_sheet) != 0.0 else self._const_expr(node.args[2], current_sheet)
            if name == "OR":
                return 1.0 if any(self._const_expr(arg, current_sheet) != 0.0 for arg in node.args) else 0.0
            if name == "IFERROR":
                return self._const_expr(node.args[0], current_sheet)
            if name in ("MIN", "MAX"):
                parts = [self._const_expr(arg, current_sheet) for arg in node.args]
                result = parts[0]
                for part in parts[1:]:
                    result = max(result, part) if name == "MAX" else min(result, part)
                return result
            if name == "SUM":
                parts = []
                for arg in node.args:
                    if arg.kind == "func" and arg.value == "XLOOKUP":
                        lookup_value, lookup_array, return_array, default = arg.args
                        lookup_sheet, lookup_address = self.split_ref(str(lookup_array.value), current_sheet)
                        return_sheet, return_address = self.split_ref(str(return_array.value), current_sheet)
                        if lookup_sheet != "增益" or return_sheet != "增益":
                            raise ValueError("常量求值仅支持增益表XLOOKUP")
                        return_col = column_index_from_string(return_address.replace("$", "").split(":", 1)[0])
                        default_value = self._const_expr(default, current_sheet)
                        parts.extend(self._const_gain(cell_id, return_col, default_value, current_sheet) for cell_id in self.range_ids(str(lookup_value.value), current_sheet))
                        continue
                    if arg.kind == "range" and ":" in str(arg.value):
                        for cell_id in self.range_ids(str(arg.value), current_sheet):
                            value = self.cell_values[cell_id]
                            if isinstance(value, (int, float)) and not isinstance(value, bool) or isinstance(value, str) and value.startswith("="):
                                parts.append(self._const_value_cell(cell_id))
                    else:
                        parts.append(self._const_expr(arg, current_sheet))
                return sum(parts)
            if name == "VLOOKUP":
                if len(node.args) != 4 or node.args[1].kind != "range" or node.args[2].kind != "number" or node.args[3].kind != "number" or node.args[3].value != 0:
                    raise ValueError("仅支持精确匹配VLOOKUP")
                key = self._const_expr(node.args[0], current_sheet)
                col = int(node.args[2].value)
                table_sheet, table_address = self.split_ref(str(node.args[1].value), current_sheet)
                if table_sheet == "武学奇术" and table_address.replace("$", "") == "A:ZZ":
                    for key_id, row in self.lookup_rows:
                        if float(key_id) == key:
                            value_id = self.cells.get(("武学奇术", self.sheets["武学奇术"].cell(row, col).coordinate))
                            return 0.0 if value_id is None else self._const_value_cell(value_id)
                    return 0.0
                pairs = self._vlookup_pairs(node, table_sheet, table_address, col, current_sheet)
                for key_id, value_id in pairs:
                    if key_id is not None and self._const_value_cell(key_id) == key:
                        return 0.0 if value_id is None else self._const_value_cell(value_id)
                return 0.0
        raise ValueError(f"常量求值失败：{node}")

    def _const_gain(self, cell_id: int, return_col: int, default_value: float, current_sheet: str) -> float:
        key = self._const_value_cell(cell_id)
        for match_id, row in self.gain_rows:
            if float(match_id) == key:
                value_id = self.cells.get(("增益", self.sheets["增益"].cell(row, return_col).coordinate))
                return default_value if value_id is None else self._const_value_cell(value_id)
        return default_value

    def _vlookup_pairs(self, node: Node, table_sheet: str, table_address: str, col: int, current_sheet: str):
        left, right = table_address.replace("$", "").split(":", 1)
        lcol, lrow = coordinate_from_string(left)
        rcol, rrow = coordinate_from_string(right)
        first_col = column_index_from_string(lcol)
        pairs = []
        for row in range(lrow, rrow + 1):
            key_id = self.cells.get((table_sheet, self.sheets[table_sheet].cell(row, first_col).coordinate))
            value_id = self.cells.get((table_sheet, self.sheets[table_sheet].cell(row, first_col + col - 1).coordinate))
            pairs.append((key_id, value_id))
        self._pairs_key_cells = {key_id for key_id, _ in pairs if key_id is not None}
        return pairs

    def defaults(self):
        result = []
        for cell_id, kind in zip(self.input_ids, self.input_kinds):
            value = 0 if cell_id is None else self.cell_values[cell_id]
            if kind == "str":
                result.append(str(value or ""))
            elif isinstance(value, (int, float)):
                result.append(float(value))
            else:
                cached = 0 if cell_id is None else self.cached_values.get(self.cell_names[cell_id])
                result.append(float(cached or 0))
        return result

    def cached(self, sheet: str, address: str) -> float:
        value = self.cached_values.get(f"{sheet}!{address}")
        if not isinstance(value, (int, float)) or not math.isfinite(float(value)):
            raise ValueError(f"{self.path.name} {sheet}!{address} 没有有效缓存值：{value!r}")
        return float(value)

    def rust(self, index: int) -> str:
        reachable, global_lookup_cols, gain_lookup_cols = self.direct_dependencies()
        order, dependencies = self.direct_order(reachable)
        input_initializers = []
        input_updates = []
        for input_index, cell_id in enumerate(self.input_ids):
            if cell_id is not None:
                input_initializers.append(
                    f"        let value = input.get({input_index}).copied().unwrap_or(0.0);\n"
                    f"        state.input_bits[{input_index}] = value.to_bits();\n"
                    f"        state.v[{cell_id}] = value;\n"
                    f"        state.dirty[{cell_id}] = epoch;"
                )
                input_updates.append(
                    f"        let value = input.get({input_index}).copied().unwrap_or(0.0);\n"
                    f"        let bits = value.to_bits();\n"
                    f"        if state.input_bits[{input_index}] != bits {{\n"
                    f"            state.input_bits[{input_index}] = bits;\n"
                    f"            state.v[{cell_id}] = value;\n"
                    f"            state.dirty[{cell_id}] = epoch;\n"
                    f"            changed_mask |= 1_u64 << {input_index};\n"
                    f"        }}"
                )
        lookup_arms = []
        for key_id, row in self.lookup_rows:
            column_arms = []
            for col in sorted(global_lookup_cols):
                cell_id = self.cells.get(("武学奇术", self.sheets["武学奇术"].cell(row, col).coordinate))
                column_arms.append(f"                    {col} => {'0.0' if cell_id is None else f'v[{cell_id}]'},")
            lookup_arms.append(f"            x if x == {f64_literal(float(key_id))} => match col {{\n" + "\n".join(column_arms) + "\n                    _ => 0.0,\n                },")
        output_ids = [self.resolve_single(ref, "期望") for ref in ("I10", "I12", "I16", "I14")]
        if any(cell_id is None for cell_id in output_ids):
            raise ValueError(f"{self.path.name} 缺少标准输出单元格")
        rd_baseline = self.cached("RD", "I14")
        gain_arms = []
        for key_id, row in self.gain_rows:
            column_arms = []
            for col in sorted(gain_lookup_cols):
                cell_id = self.cells.get(("增益", self.sheets["增益"].cell(row, col).coordinate))
                column_arms.append(f"                    {col} => {'default' if cell_id is None else f'v[{cell_id}]'},")
            gain_arms.append(f"            x if x == {f64_literal(float(key_id))} => match col {{\n" + "\n".join(column_arms) + "\n                    _ => default,\n                },")
        input_ids = {cell_id for cell_id in self.input_ids if cell_id is not None}
        dynamic = set(input_ids)
        for cell_id in order:
            if any(dependency in dynamic for dependency in dependencies[cell_id]):
                dynamic.add(cell_id)
        self._const_ids = set(order) - input_ids - dynamic
        self._const_cache = {}
        for cell_id in order:
            if cell_id in self._const_ids:
                self._const_cache[cell_id] = self._const_value_of(cell_id)
        constant_assignments = [f"    v[{cell_id}] = {self.direct_value_expr(cell_id)};" for cell_id in order if cell_id not in input_ids and cell_id not in dynamic]
        input_masks = [0 for _ in self.cell_values]
        for input_index, cell_id in enumerate(self.input_ids):
            if cell_id is not None:
                input_masks[cell_id] |= 1 << input_index
        for cell_id in order:
            for dependency in dependencies[cell_id]:
                input_masks[cell_id] |= input_masks[dependency]
        dynamic_assignments = []
        for cell_id in order:
            if cell_id in input_ids or cell_id not in dynamic:
                continue
            dynamic_dependencies = sorted(dependency for dependency in dependencies[cell_id] if dependency in dynamic)
            if not dynamic_dependencies:
                raise ValueError(f"{self.path.name} 动态单元格缺少动态依赖：{self.cell_names[cell_id]}")
            dirty_condition = " || ".join(f"dirty[{dependency}] == epoch" for dependency in dynamic_dependencies)
            dynamic_assignments.append((
                f"    if {dirty_condition} {{\n"
                f"        let old_bits = v[{cell_id}].to_bits();\n"
                f"        v[{cell_id}] = {self.direct_value_expr(cell_id)};\n"
                f"        if v[{cell_id}].to_bits() != old_bits {{ dirty[{cell_id}] = epoch; }}\n"
                f"    }}",
                input_masks[cell_id],
            ))
        def make_chunks(prefix, assignments):
            result = []
            for chunk_index in range(0, len(assignments), 192):
                name = f"{prefix}_{chunk_index // 192}"
                result.append((name, assignments[chunk_index:chunk_index + 192]))
            return result
        constant_chunks = make_chunks(f"constant_{index}_chunk", constant_assignments)
        dynamic_chunks = []
        for chunk_index in range(0, len(dynamic_assignments), 192):
            entries = dynamic_assignments[chunk_index:chunk_index + 192]
            name = f"calculate_{index}_chunk_{chunk_index // 192}"
            mask = 0
            for _, entry_mask in entries:
                mask |= entry_mask
            dynamic_chunks.append((name, [line for line, _ in entries], mask))
        chunks = constant_chunks + [(name, lines) for name, lines, _ in dynamic_chunks]
        chunk_functions = "\n".join(
            f"#[inline({'never' if name.startswith('constant_') else 'always'})]\n"
            f"fn {name}(v: &mut [f64; {len(self.cell_values)}]"
            f"{'' if name.startswith('constant_') else f', dirty: &mut [u32; {len(self.cell_values)}], epoch: u32'}) {{\n"
            f"{chr(10).join(lines)}\n}}"
            for name, lines in chunks
        )
        constant_calls = "\n".join(f"        {name}(&mut v);" for name, _ in constant_chunks)
        dynamic_groups = []
        for group_index in range(0, len(dynamic_chunks), 4):
            name = f"calculate_{index}_group_{group_index // 4}"
            calls = "\n".join(
                f"    if changed_mask & 0x{mask:016x} != 0 {{ {chunk_name}(v, dirty, epoch); }}"
                for chunk_name, _, mask in dynamic_chunks[group_index:group_index + 4]
            )
            dynamic_groups.append((name, calls))
        group_functions = "\n".join(f"#[inline(never)]\nfn {name}(v: &mut [f64; {len(self.cell_values)}], dirty: &mut [u32; {len(self.cell_values)}], epoch: u32, changed_mask: u64) {{\n{calls}\n}}" for name, calls in dynamic_groups)
        dynamic_calls = "\n".join(f"        {name}(&mut state.v, &mut state.dirty, epoch, changed_mask);" for name, _ in dynamic_groups)
        batch_assignments = [f"    v[{cell_id}] = {self.direct_value_expr(cell_id)};" for cell_id in order if cell_id not in input_ids and cell_id in dynamic]
        batch_chunks = make_chunks(f"batch_{index}_chunk", batch_assignments)
        batch_chunk_functions = "\n".join(
            f"#[inline(always)]\nfn {name}(v: &mut [f64; {len(self.cell_values)}]) {{\n{chr(10).join(lines)}\n}}"
            for name, lines in batch_chunks
        )
        batch_dynamic_calls = "\n".join(f"        {name}(&mut **state);" for name, _ in batch_chunks)
        batch_input_sets = "\n".join(
            f"        state[{cell_id}] = input[offset + {input_index}];"
            for input_index, cell_id in enumerate(self.input_ids) if cell_id is not None
        )
        return f"""
fn vlookup(key: f64, col: usize, v: &[f64; {len(self.cell_values)}]) -> f64 {{
        match key {{
{chr(10).join(lookup_arms)}
            _ => 0.0,
        }}
}}
fn local_vlookup(key: f64, pairs: &[(usize, usize)], v: &[f64; {len(self.cell_values)}]) -> f64 {{
        for &(key_id, value_id) in pairs {{
            if key_id != usize::MAX && v[key_id] == key {{
                return if value_id == usize::MAX {{ 0.0 }} else {{ v[value_id] }};
            }}
        }}
        0.0
}}
fn xlookup_gain(key: f64, col: usize, default: f64, v: &[f64; {len(self.cell_values)}]) -> f64 {{
        match key {{
{chr(10).join(gain_arms)}
            _ => default,
        }}
}}

{chunk_functions}
{group_functions}
{batch_chunk_functions}

struct EngineState {{
    v: Box<[f64; {len(self.cell_values)}]>,
    dirty: Box<[u32; {len(self.cell_values)}]>,
    input_bits: [u64; {INPUT_LEN}],
    epoch: u32,
    initialized: bool,
}}
struct EngineWorkspace(std::cell::UnsafeCell<EngineState>);
unsafe impl Sync for EngineWorkspace {{}}
struct BatchWorkspace(std::cell::UnsafeCell<Box<[f64; {len(self.cell_values)}]>>);
unsafe impl Sync for BatchWorkspace {{}}

fn get_workspace() -> &'static mut EngineState {{
    static WORKSPACE: std::sync::OnceLock<EngineWorkspace> = std::sync::OnceLock::new();
    let workspace = WORKSPACE.get_or_init(|| {{
        let mut v: Box<[f64; {len(self.cell_values)}]> = vec![0.0; {len(self.cell_values)}].into_boxed_slice().try_into().ok().unwrap();
{constant_calls}
        EngineWorkspace(std::cell::UnsafeCell::new(EngineState {{
            v,
            dirty: vec![0_u32; {len(self.cell_values)}].into_boxed_slice().try_into().ok().unwrap(),
            input_bits: [0_u64; {INPUT_LEN}],
            epoch: 0,
            initialized: false,
        }}))
    }});
    // A WebAssembly module instance is single-threaded and this ABI is deliberately non-reentrant.
    unsafe {{ &mut *workspace.0.get() }}
}}

fn calculate_{index}(input: &[f64], output: &mut [f64]) {{
    let state = get_workspace();
    state.epoch = state.epoch.wrapping_add(1);
    if state.epoch == 0 {{
        state.dirty.fill(0);
        state.epoch = 1;
    }}
    let epoch = state.epoch;
    let mut changed_mask = 0_u64;
    if !state.initialized {{
        // The first call must evaluate the complete dynamic graph. A zero-valued
        // intermediate can still feed a non-zero formula together with constants.
        state.dirty.fill(epoch);
{chr(10).join(input_initializers)}
        changed_mask = (1_u64 << {INPUT_LEN}) - 1;
        state.initialized = true;
    }} else {{
{chr(10).join(input_updates)}
    }}
    if changed_mask != 0 {{
{dynamic_calls}
    }}
    let total = state.v[{output_ids[0]}];
    let adps = state.v[{output_ids[1]}];
    let graduation = state.v[{output_ids[2]}];
    let rdps = state.v[{output_ids[3]}];
    output[0] = total;
    output[1] = adps;
    output[2] = graduation;
    output[3] = rdps;
    output[4] = rdps / {f64_literal(rd_baseline)};
}}

#[unsafe(no_mangle)]
pub unsafe extern "C" fn yysls_calc_batch(input: *const f64, count: usize, output: *mut f64) {{
    let input = unsafe {{ slice::from_raw_parts(input, INPUT_LEN * count) }};
    let output = unsafe {{ slice::from_raw_parts_mut(output, OUTPUT_LEN * count) }};
static BATCH_WORKSPACE: std::sync::OnceLock<BatchWorkspace> = std::sync::OnceLock::new();
    let workspace = BATCH_WORKSPACE.get_or_init(|| {{
        let mut v: Box<[f64; {len(self.cell_values)}]> = vec![0.0; {len(self.cell_values)}].into_boxed_slice().try_into().ok().unwrap();
{constant_calls}
        BatchWorkspace(std::cell::UnsafeCell::new(v))
    }});
    let state = unsafe {{ &mut *workspace.0.get() }};
    for i in 0..count {{
        let offset = i * INPUT_LEN;
{batch_input_sets}
{batch_dynamic_calls}
        output[i * OUTPUT_LEN + 0] = state[{output_ids[0]}];
        output[i * OUTPUT_LEN + 1] = state[{output_ids[1]}];
        output[i * OUTPUT_LEN + 2] = state[{output_ids[2]}];
        output[i * OUTPUT_LEN + 3] = state[{output_ids[3]}];
        output[i * OUTPUT_LEN + 4] = state[{output_ids[3]}] / {f64_literal(rd_baseline)};
    }}
}}
"""

    def standalone_rust(self) -> str:
        body = self.rust(0)
        return """use std::slice;
const INPUT_LEN: usize = 40;
const OUTPUT_LEN: usize = 5;
""" + body + """
#[unsafe(no_mangle)]
pub extern "C" fn yysls_alloc_f64(len: usize) -> *mut f64 {
    Box::into_raw(vec![0.0_f64; len].into_boxed_slice()).cast::<f64>()
}
#[unsafe(no_mangle)]
pub unsafe extern "C" fn yysls_free_f64(ptr: *mut f64, len: usize) {
    if !ptr.is_null() { drop(unsafe { Box::from_raw(std::ptr::slice_from_raw_parts_mut(ptr, len)) }); }
}
#[unsafe(no_mangle)]
pub extern "C" fn yysls_class_input_len() -> i32 { INPUT_LEN as i32 }
#[unsafe(no_mangle)]
pub extern "C" fn yysls_class_output_len() -> i32 { OUTPUT_LEN as i32 }
#[unsafe(no_mangle)]
pub unsafe extern "C" fn yysls_calc_class(_flow_id: i32, input: *const f64) -> f64 {
    let input = unsafe { slice::from_raw_parts(input, INPUT_LEN) };
    let mut output = [0.0; OUTPUT_LEN];
    calculate_0(input, &mut output);
    output[0]
}
#[unsafe(no_mangle)]
pub unsafe extern "C" fn yysls_calc_class_outputs(_flow_id: i32, input: *const f64, output: *mut f64) {
    let input = unsafe { slice::from_raw_parts(input, INPUT_LEN) };
    let output = unsafe { slice::from_raw_parts_mut(output, OUTPUT_LEN) };
    calculate_0(input, output);
}
"""

def workbook_specs():
    files = sorted(EXCEL_DIR.glob("*.xlsx"))
    specs = []
    for class_name in FLOW_ORDER:
        matches = [path for path in files if path.name.startswith(class_name + "110阶")]
        if class_name == "牵丝翊":
            matches.sort(key=lambda path: ("2.0" not in path.stem, path.name))
        if not matches:
            raise ValueError(f"缺少流派Excel：{class_name}")
        for path in matches:
            version_match = re.search(r"计算器([0-9.]+)\.xlsx$", path.name)
            if not version_match:
                raise ValueError(f"文件名缺少版本号：{path.name}")
            version = version_match.group(1)
            flow_name = f"{class_name}@{version}" if class_name == "牵丝翊" else class_name
            specs.append((class_name, flow_name, version, path))
    if len(specs) != 11:
        raise ValueError(f"必须恰好有11份Excel，实际为{len(specs)}")
    return specs


specs = workbook_specs()
records = []
DIRECT_SOURCE_DIR.mkdir(parents=True, exist_ok=True)
for generated_source in DIRECT_SOURCE_DIR.glob("excel_*.rs"):
    generated_source.unlink()
for class_name, flow_name, version, path in specs:
    base_key = class_name if class_name in source_metadata["flowClassFields"] else next(
        key for key, mapped in source_metadata.get("flowClassNames", {}).items() if mapped == class_name
    )
    base_fields = source_metadata["flowClassFields"][base_key]
    base_kinds = source_metadata["flowClassKinds"][base_key]
    base_cells = list(source_metadata["flowClassCells"][base_key])
    # 破竹鸢 2.4 将第三、第四心法输入由 C22/C24 移到了 E22/E24。
    # 字段名称保持公开 API 兼容，但单元格位置必须以当前工作簿为准。
    if class_name == "破竹鸢" and version == "2.4":
        base_cells[base_fields.index("third_xinfa")] = "期望!E22"
        base_cells[base_fields.index("fourth_xinfa")] = "期望!E24"
    compiler = WorkbookCompiler(path, flow_name, class_name, base_cells, base_kinds)
    module_name = f"excel_{FLOW_SLUGS[class_name]}_{version.replace('.', '_')}.wasm"
    (DIRECT_SOURCE_DIR / module_name.replace(".wasm", ".rs")).write_text(compiler.standalone_rust(), encoding="utf-8")
    records.append((class_name, flow_name, version, path, module_name, list(base_fields), list(base_kinds),
                    [compiler.cell_names[cell_id] if cell_id is not None else base_cells[i]
                     for i, cell_id in enumerate(compiler.input_ids)],
                    compiler.defaults(), compiler.workbook.properties.modified,
                    compiler.cached("期望", "I10"), compiler.cached("期望", "I8"),
                    compiler.cached("期望", "I12"), compiler.cached("期望", "I14")))
    compiler.workbook.close()
    del compiler
    gc.collect()
for key in ("flowIds", "flowKeys", "flowClassNames", "flowClassFields", "flowClassKinds", "flowClassCells", "flowClassDefaultValues", "flowGraduationProfiles", "classRotationStats"):
    metadata[key] = {}
metadata["flowNames"] = []
metadata["classTableVersions"] = {"牵丝翊": []}
metadata["flowExcelModules"] = {}

manifest = {"inputLength": INPUT_LEN, "outputLength": 5, "workbooks": []}
for index, record in enumerate(records):
    (class_name, flow_name, version, path, module_name, fields, kinds, cells, defaults, modified,
     baseline_total, use_time, baseline_dps, baseline_rdps) = record
    update_time = modified.replace(tzinfo=timezone.utc).astimezone().strftime("%Y年%-m月%-d日 %H:%M:%S") if modified else ""
    metadata["flowIds"][flow_name] = index
    metadata["flowKeys"][flow_name] = f"excel_{index}"
    metadata["flowClassNames"][flow_name] = class_name
    metadata["flowNames"].append(flow_name)
    metadata["flowClassFields"][flow_name] = fields
    metadata["flowClassKinds"][flow_name] = kinds
    metadata["flowClassCells"][flow_name] = cells
    metadata["flowClassDefaultValues"][flow_name] = defaults
    metadata["flowGraduationProfiles"][flow_name] = {
        "fields": fields, "values": defaults, "workbook": path.name, "version": version,
    }
    metadata["flowExcelModules"][flow_name] = {"file": module_name, "version": version}
    metadata["classRotationStats"][flow_name] = {
        "baseline": baseline_total, "baselineTotal": baseline_total,
        "useTime": use_time, "dps": baseline_dps, "baselineDps": baseline_dps,
        "baselineRdps": baseline_rdps, "version": path.stem, "updateTime": update_time,
    }
    if class_name != "牵丝翊":
        metadata["classRotationStats"][class_name] = metadata["classRotationStats"][flow_name]
    else:
        metadata["classTableVersions"][class_name].append({
            "key": version, "label": version, "flowName": flow_name, "default": version == "2.0"
        })
    manifest["workbooks"].append({
        "id": index, "className": class_name, "flowName": flow_name, "version": version,
        "file": path.name, "module": module_name, "inputs": len(fields), "baselineTotal": baseline_total,
        "baselineDps": baseline_dps, "baselineRdps": baseline_rdps, "useTime": use_time,
    })

metadata["classDefaultValues"] = {name: metadata["flowClassDefaultValues"][name] for name in FLOW_ORDER if name in metadata["flowClassDefaultValues"]}
metadata["classDefaultValues"]["牵丝翊"] = metadata["flowClassDefaultValues"]["牵丝翊@2.0"]
if site_update_time := os.environ.get("YYSLS_SITE_UPDATE_TIME"):
    metadata["siteUpdateTime"] = site_update_time
METADATA_PATH.write_text("// Generated from calculator workbooks. Do not edit by hand.\nwindow.YYSLS_CALC_METADATA=" + json.dumps(metadata, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
STRINGS_PATH.write_text("// Generated from calculator workbooks. Do not edit by hand.\nwindow.YYSLS_CALC_STRINGS=" + json.dumps(strings, ensure_ascii=False, separators=(",", ":")) + ";\nwindow.YYSLS_CALC_STRING_IDS=" + json.dumps(string_ids, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"status": "ok", "workbooks": len(records), "strings": len(strings)}, ensure_ascii=False))
