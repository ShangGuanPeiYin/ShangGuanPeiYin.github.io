#!/usr/bin/env python3
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
EXCEL_DIR = ROOT / "static/tools/yysls-tiaolv/excels"
EXPECTED = {"牵丝玉": 1, "牵丝翊": 2, "牵丝霖": 1, "破竹尘": 1, "破竹风": 1, "破竹鸢": 1, "裂石威": 1, "裂石钧": 1, "鸣金虹": 1, "鸣金影": 1}
ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}

files = sorted(EXCEL_DIR.glob("*.xlsx"))
counts = {name: 0 for name in EXPECTED}
versions = set()
problems = []
for path in files:
    match = re.match(r"(.+?)110阶.*计算器([0-9.]+)\.xlsx$", path.name)
    if not match or match.group(1) not in EXPECTED:
        problems.append(f"无法识别工作簿：{path.name}")
        continue
    flow, version = match.groups()
    key = (flow, version)
    if key in versions:
        problems.append(f"重复版本：{flow} {version}")
    versions.add(key)
    counts[flow] += 1
    formulas = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=True)
    if getattr(formulas, "_external_links", None):
        problems.append(f"存在外部链接：{path.name}")
    for sheet in cached.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.upper() in ERRORS:
                    problems.append(f"缓存公式错误：{path.name} {sheet.title}!{cell.coordinate}={cell.value}")
    formulas.close()
    cached.close()

if len(files) != 11:
    problems.append(f"工作簿数量应为11，实际为{len(files)}")
for flow, expected in EXPECTED.items():
    if counts[flow] != expected:
        problems.append(f"{flow}应有{expected}份，实际为{counts[flow]}")
if problems:
    raise SystemExit("\n".join(problems))
print(json.dumps({"status": "ok", "workbooks": len(files), "flows": counts}, ensure_ascii=False))
